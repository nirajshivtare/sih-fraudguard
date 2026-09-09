import os
from pptx import Presentation

def replace_text_in_shape(shape, old_text, new_text):
    if not shape.has_text_frame:
        return
    for paragraph in shape.text_frame.paragraphs:
        if old_text in paragraph.text:
            paragraph.text = paragraph.text.replace(old_text, new_text)

def build_ppt(template_path, output_path):
    prs = Presentation(template_path)

    # --- Slide 1: TITLE PAGE ---
    slide1 = prs.slides[0]
    # In slide 1 we have specific text to replace
    for shape in slide1.shapes:
        if shape.has_text_frame:
            text = shape.text
            if "Problem Statement ID" in text:
                shape.text_frame.clear()
                p = shape.text_frame.paragraphs[0]
                p.text = "Problem Statement ID - PS 106\nProblem Statement Title - AI-Powered Email Threat Detection, GeoLocation and Forensic Intelligence Platform\nTheme - Cybersecurity\nPS Category - Software\nTeam ID - [Your Team ID]\nTeam Name - [Your Team Name]"

    # --- Slide 2: IDEA TITLE ---
    slide2 = prs.slides[1]
    for shape in slide2.shapes:
        if shape.has_text_frame:
            if "IDEA TITLE" in shape.text:
                shape.text_frame.clear()
                p = shape.text_frame.paragraphs[0]
                p.text = "IDEA TITLE: AI-Powered Email Threat Detection & Forensic Intelligence"
            elif "Proposed Solution" in shape.text:
                shape.text_frame.clear()
                p = shape.text_frame.paragraphs[0]
                p.text = "Proposed Solution:\n• AI-driven forensic platform to detect threats & trace origins.\n• Extracts headers, reconstructs relay paths, analyzes SPF/DKIM/DMARC.\n\nHow it addresses the problem:\n• Goes beyond basic blocking by providing deep forensic tracing.\n• Uncovers true source infrastructure using IP/Domain intelligence.\n\nInnovation and uniqueness:\n• Combines NLP/ML for text analysis with Neo4j Graph DB for campaign attribution."

    # --- Slide 3: TECHNICAL APPROACH ---
    slide3 = prs.slides[2]
    for shape in slide3.shapes:
        if shape.has_text_frame:
            if "Technologies to be used" in shape.text:
                shape.text_frame.clear()
                p = shape.text_frame.paragraphs[0]
                p.text = "Technologies to be used:\n• Backend: Python (FastAPI), React.js (Frontend Dashboard)\n• AI/ML: HuggingFace Transformers, spaCy, Scikit-learn\n• Forensics/DB: PostgreSQL, Neo4j (Graph), MaxMind GeoIP\n\nMethodology & Process:\n1. Ingestion: Parse .eml files & validate auth protocols.\n2. AI Analysis: ML classifies phishing & impersonation language.\n3. Traceability: Extract IP hops & map GeoLocation.\n4. Correlation: Map domain/IP relationships to identify campaigns.\n5. Reporting: Dashboard alerts & PDF forensic reports."

    # --- Slide 4: FEASIBILITY AND VIABILITY ---
    slide4 = prs.slides[3]
    for shape in slide4.shapes:
        if shape.has_text_frame:
            if "Analysis of the feasibility" in shape.text:
                shape.text_frame.clear()
                p = shape.text_frame.paragraphs[0]
                p.text = "Feasibility Analysis:\n• High feasibility utilizing robust open-source ML libraries and scalable frameworks (FastAPI).\n• IP and Threat Intel APIs are well-documented and readily available.\n\nPotential Challenges & Risks:\n• Processing large volumes of emails with low latency.\n• Ensuring Data Privacy (PII) during content analysis.\n\nStrategies for Overcoming:\n• Implement async processing and caching.\n• Use data masking/anonymization before ML processing."

    # --- Slide 5: IMPACT AND BENEFITS ---
    slide5 = prs.slides[4]
    for shape in slide5.shapes:
        if shape.has_text_frame:
            if "Potential impact on the target audience" in shape.text:
                shape.text_frame.clear()
                p = shape.text_frame.paragraphs[0]
                p.text = "Potential Impact:\n• Cyber Response Teams: Drastically reduces investigation time.\n• Organizations: Secures ecosystems against targeted BEC attacks.\n\nBenefits of the Solution:\n• Financial: Prevents major losses from payment diversion.\n• Security: Real-time blocking of AI-generated phishing.\n• Legal: Provides chain-of-custody reporting for law enforcement."

    # --- Slide 6: RESEARCH AND REFERENCES ---
    slide6 = prs.slides[5]
    for shape in slide6.shapes:
        if shape.has_text_frame:
            if "Details / Links" in shape.text:
                shape.text_frame.clear()
                p = shape.text_frame.paragraphs[0]
                p.text = "References & Research:\n• FBI IC3 Reports on Business Email Compromise (BEC).\n• MITRE ATT&CK Framework: Phishing and Initial Access techniques.\n• Documentation on Email Authentication: SPF, DKIM, DMARC RFCs.\n• Graph-based Threat Intelligence correlation research."

    # Delete slide 7
    xml_slides = prs.slides._sldIdLst  
    slides = list(xml_slides)
    if len(slides) > 6:
        xml_slides.remove(slides[6])

    prs.save(output_path)
    print(f"Filled PPT saved to: {output_path}")

if __name__ == "__main__":
    template = r"C:\Users\niraj\Downloads\SIH2026-IDEA-Presentation-Format.pptx"
    output = r"C:\Users\niraj\OneDrive\Desktop\Antigravity\SIH 2026\SIH_2026_PS106_Filled.pptx"
    build_ppt(template, output)
