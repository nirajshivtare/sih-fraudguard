from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

prs = Presentation()

def add_slide(prs, layout_index, title_text, text_lines):
    slide = prs.slides.add_slide(prs.slide_layouts[layout_index])
    title = slide.shapes.title
    title.text = title_text
    
    body_shape = slide.placeholders[1]
    tf = body_shape.text_frame
    
    for i, line in enumerate(text_lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(20)

# Slide 1: Title (Matches standard SIH Title Slide)
slide = prs.slides.add_slide(prs.slide_layouts[0])
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "PS-106: AI-Powered Email Threat Detection & Forensic Intelligence"
subtitle.text = "Team Name: [Your Team Name]\nMinistry/Organization: [Specific Ministry]\nSmart India Hackathon 2026"

# Slide 2: Idea / Approach Details
add_slide(
    prs, 1,
    "Idea / Approach Details",
    [
        "Problem: Sophisticated spoofing and fraud evade standard email filters.",
        "Proposed Solution: An AI-driven forensic platform to detect threats and trace origins.",
        "• AI Detection Engine: NLP & ML classify phishing, BEC, and urgency cues.",
        "• Header & Protocol Forensics: Deep analysis of SPF/DKIM/DMARC and routing.",
        "• Origin Traceability: Extracts IP hop chains, maps GeoLocation, and WHOIS data.",
        "• Graph Correlation: Maps relationships between domains, IPs, and campaigns."
    ]
)

# Slide 3: Use Cases
add_slide(
    prs, 1,
    "Use Cases",
    [
        "1. Institutional Security: Real-time blocking of sophisticated phishing before user interaction.",
        "2. Forensic Investigation: Assisting law enforcement with visual trace maps of attacker origins.",
        "3. Financial Fraud Prevention: Flagging compromised accounts and invoice diversion attempts.",
        "4. Threat Intelligence Sharing: Correlating campaigns across domains to prevent repeat attacks."
    ]
)

# Slide 4: Dependencies / Show Stoppers
add_slide(
    prs, 1,
    "Dependencies / Show Stoppers",
    [
        "Dependencies:",
        "• Access to diverse datasets of phishing and legitimate emails to train the ML models.",
        "• Reliable third-party APIs for IP GeoLocation and Domain/WHOIS Intelligence.",
        "Show Stoppers (Risks):",
        "• Data Privacy (PII): Analyzing email content must comply with strict privacy regulations.",
        "• Processing Latency: Deep ML analysis could delay real-time email delivery if not optimized."
    ]
)

# Slide 5: Tech Stack
add_slide(
    prs, 1,
    "Tech Stack",
    [
        "Backend / Processing:",
        "• Python (FastAPI): High performance, asynchronous API handling.",
        "AI / ML Engine:",
        "• HuggingFace Transformers, spaCy, Scikit-learn.",
        "Database & Visualization:",
        "• Neo4j (Graph DB for attribution), PostgreSQL.",
        "Frontend:",
        "• React.js / Tailwind CSS for the Analyst Dashboard."
    ]
)

# Slide 6: Team Member Details
add_slide(
    prs, 1,
    "Team Member Details",
    [
        "1. [Name] - Team Leader - [Domain/Skill]",
        "2. [Name] - Member - [Domain/Skill]",
        "3. [Name] - Member - [Domain/Skill]",
        "4. [Name] - Member - [Domain/Skill]",
        "5. [Name] - Member - [Domain/Skill]",
        "6. [Name] - Member - [Domain/Skill]"
    ]
)

ppt_filename = "SIH_2026_Standard_Format.pptx"
prs.save(ppt_filename)
