from pptx import Presentation
import sys

def read_ppt(file_path):
    try:
        prs = Presentation(file_path)
        for i, slide in enumerate(prs.slides):
            print(f"--- Slide {i+1} ---")
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    print(shape.text)
            print("\n")
    except Exception as e:
        print(f"Error reading pptx: {e}")

if __name__ == "__main__":
    read_ppt(r"C:\Users\niraj\Downloads\SIH2026-IDEA-Presentation-Format.pptx")
