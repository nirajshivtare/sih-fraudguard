import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()

# Function to add a title slide
def add_title_slide(prs, title_text, subtitle_text):
    slide_layout = prs.slide_layouts[0] # 0 is title slide layout
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = title_text
    subtitle.text = subtitle_text

# Function to add a slide with title and content
def add_content_slide(prs, title_text, content_list):
    slide_layout = prs.slide_layouts[1] # 1 is title and content layout
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body_shape = slide.shapes.placeholders[1]
    
    title.text = title_text
    tf = body_shape.text_frame
    
    for i, item in enumerate(content_list):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(22)

# --- Add Slides ---

# Slide 1: Title
add_title_slide(
    prs, 
    "AI-Powered Email Threat Detection, GeoLocation\nand Forensic Intelligence Platform", 
    "Smart India Hackathon (SIH) 2026\nProblem Statement: PS 106"
)

# Slide 2: Background
add_content_slide(
    prs,
    "Background: The Phishing Epidemic",
    [
        "Email remains a widely used, yet highly exploited communication channel.",
        "Attack vectors: Phishing, impersonation, BEC (Business Email Compromise), fraud, and malware.",
        "Threat actors use spoofed domains, deceptive identities, and AI-generated text.",
        "Traditional spam filters and rule-based systems are insufficient against modern attacks."
    ]
)

# Slide 3: The Problem Statement
add_content_slide(
    prs,
    "The Problem",
    [
        "Current tools block suspicious content but lack deep forensic tracing.",
        "They fail to correlate email headers, relay paths, SPF/DKIM/DMARC, IP reputation, and geolocation.",
        "Organizations lack the capability to trace the source path and identify sender infrastructure.",
        "A gap exists in generating forensic intelligence to identify the attacker's true origin."
    ]
)

# Slide 4: Proposed Solution
add_content_slide(
    prs,
    "Our Proposed Solution",
    [
        "An AI-Powered Platform combining NLP, ML, and Forensic Intelligence.",
        "Ingests raw emails and validates sender authentication.",
        "Extracts IOCs (Indicators of Compromise) and reconstructs relay paths.",
        "Analyzes originating IPs and provides GeoLocation data.",
        "Generates confidence-based fraud assessments and visual trace maps."
    ]
)

# Slide 5: Key Components (Part 1)
add_content_slide(
    prs,
    "Key Components: Detection & Analysis",
    [
        "1. Fraudulent Email Detection Engine: NLP-based analysis of subject/body, detecting phishing, spoofing, and BEC patterns.",
        "2. Email Header & Protocol Analysis: Deep analysis of Return-Path, DKIM, SPF, DMARC, and mail routing anomalies.",
        "3. Origin Traceability & Location Analysis: IP extraction, GeoLocation mapping, and Domain Intelligence (WHOIS, DNS)."
    ]
)

# Slide 6: Key Components (Part 2)
add_content_slide(
    prs,
    "Key Components: Correlation & Reporting",
    [
        "4. Identity Correlation & Attribution: Threat intelligence correlation and Graph-based relationship analysis.",
        "5. Alerting & Dashboard: Real-time alerts, Analyst dashboard with fraud scores, map tracing, and visual indicators.",
        "6. Privacy & Compliance Safeguards: Secure handling of metadata, evidence preservation, and chain-of-custody support."
    ]
)

# Slide 7: Expected Outcomes
add_content_slide(
    prs,
    "Expected Outcomes",
    [
        "Early and accurate detection of fraudulent/phishing attacks.",
        "Improved capability to trace suspicious origins and infrastructure.",
        "Enhanced fraud investigation through geolocation and domain intel.",
        "Reduction in financial loss and data leaks caused by email fraud.",
        "Better institutional readiness for cyber incident response."
    ]
)

# Slide 8: Proposed Tech Stack
add_content_slide(
    prs,
    "Proposed Technology Stack",
    [
        "Backend: Python (FastAPI/Django) for high performance AI integration.",
        "Frontend: React.js for an interactive Analyst Dashboard.",
        "AI/NLP: HuggingFace Transformers, spaCy, Scikit-learn.",
        "Database: PostgreSQL (Data) + Neo4j (Graph correlation).",
        "Forensics: MaxMind GeoIP, WHOIS APIs, Threat Intel integrations."
    ]
)

# Save the presentation
ppt_filename = "SIH_2026_PS106_Presentation.pptx"
prs.save(ppt_filename)
print(f"Presentation saved successfully as {ppt_filename}")
